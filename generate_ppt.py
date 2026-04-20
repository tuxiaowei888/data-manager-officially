"""
生成数据要素服务全链路 PPT 页面
深色科技风格，与原有底板保持一致
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

# 颜色定义 - 深色科技风格
COLORS = {
    'background': RGBColor(10, 25, 50),
    'primary': RGBColor(0, 150, 255),
    'secondary': RGBColor(0, 200, 255),
    'accent': RGBColor(255, 200, 0),
    'text_white': RGBColor(255, 255, 255),
    'text_light': RGBColor(200, 220, 255),
    'text_gray': RGBColor(150, 170, 200),
    'card_bg': RGBColor(20, 40, 70),
}


def create_data_service_flow_ppt():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    add_background(slide, prs)
    add_left_sidebar(slide, prs)
    add_title(slide)
    add_flow_chart(slide)
    add_footer(slide, prs)
    
    output_path = '数据要素服务全链路.pptx'
    prs.save(output_path)
    print(f"✓ PPT 已生成：{output_path}")
    print(f"✓ 可以在 WPS 中打开并编辑")
    
    return output_path


def add_background(slide, prs):
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(0),
        prs.slide_width, prs.slide_height
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLORS['background']
    bg.line.fill.background()


def add_left_sidebar(slide, prs):
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(0),
        Inches(0.8), prs.slide_height
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLORS['primary']
    bar.line.fill.background()
    
    gradient_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.8), Inches(0),
        Inches(0.3), prs.slide_height
    )
    gradient_bar.fill.solid()
    gradient_bar.fill.fore_color.rgb = RGBColor(0, 100, 200)
    gradient_bar.line.fill.background()
    gradient_bar.fill.transparency = 0.5


def add_title(slide):
    title_box = slide.shapes.add_textbox(
        Inches(1.5), Inches(0.4),
        Inches(10), Inches(1)
    )
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "数据要素服务全链路"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = COLORS['text_white']
    p.alignment = PP_ALIGN.CENTER
    
    subtitle_box = slide.shapes.add_textbox(
        Inches(1.5), Inches(1.2),
        Inches(10), Inches(0.5)
    )
    tf = subtitle_box.text_frame
    p = tf.paragraphs[0]
    p.text = "一站式数据价值化解决方案"
    p.font.size = Pt(20)
    p.font.color.rgb = COLORS['secondary']
    p.alignment = PP_ALIGN.CENTER


def add_flow_chart(slide):
    stage1_x = Inches(2)
    stage2_x = Inches(4.5)
    stage3_x = Inches(7)
    stage_y = Inches(2.5)
    box_width = Inches(1.8)
    box_height = Inches(1)
    
    add_process_box(slide, stage1_x, stage_y, box_width, box_height, 
                   "数据盘点", "基础梳理", is_start=True)
    
    add_arrow(slide, stage1_x + box_width, stage_y + box_height/2,
              stage2_x, stage_y + box_height/2)
    
    add_process_box(slide, stage2_x, stage_y, box_width, box_height,
                   "数据治理", "质量提升")
    
    add_arrow(slide, stage2_x + box_width, stage_y + box_height/2,
              stage3_x, stage_y + box_height/2)
    
    add_process_box(slide, stage3_x, stage_y, box_width, box_height,
                   "数据合规", "合规评估", is_compliance=True)
    
    确权_y = Inches(4.3)
    确权_width = Inches(2.2)
    确权_height = Inches(1.1)
    
    add_curved_arrow(slide, stage3_x + box_width/2, stage_y + box_height,
                     stage3_x, 确权_y)
    
    add_curved_arrow(slide, stage3_x + box_width/2, stage_y + box_height,
                     stage3_x + box_width + Inches(0.5), 确权_y)
    
    add_process_box(slide, stage3_x - 确权_width/2, 确权_y, 确权_width, 确权_height,
                   "登记确权", "官方登记", is_branch=True)
    
    add_process_box(slide, stage3_x + box_width - 确权_width/2, 确权_y, 确权_width, 确权_height,
                   "协议确权", "合同约定", is_branch=True)
    
    转化_y = Inches(6)
    转化_width = Inches(1.6)
    转化_height = Inches(0.9)
    转化_spacing = Inches(0.4)
    转化_start_x = stage3_x - (转化_width * 2 + 转化_spacing)
    
    add_arrow_down(slide, stage3_x - 确权_width/4, 确权_y + 确权_height,
                   转化_start_x + 转化_width/2, 转化_y)
    
    add_arrow_down(slide, stage3_x + box_width + 确权_width/4, 确权_y + 确权_height,
                   转化_start_x + 转化_width * 1.5 + 转化_spacing/2, 转化_y)
    
    add_process_box(slide, 转化_start_x, 转化_y, 转化_width, 转化_height,
                   "融资", "质押贷款", is_final=True)
    
    add_process_box(slide, 转化_start_x + 转化_width + 转化_spacing/2, 转化_y,
                   转化_width, 转化_height, "入表", "资产入表", is_final=True)
    
    add_process_box(slide, 转化_start_x + (转化_width + 转化_spacing/2) * 2, 转化_y,
                   转化_width, 转化_height, "授权交易", "流通变现", is_final=True)
    
    add_process_box(slide, 转化_start_x + (转化_width + 转化_spacing/2) * 3, 转化_y,
                   转化_width, 转化_height, "作价入股", "资本运作", is_final=True)
    
    add_core_capability(slide)


def add_process_box(slide, left, top, width, height, title, subtitle="",
                   is_start=False, is_compliance=False, is_branch=False, is_final=False):
    
    if is_start:
        fill_color = COLORS['primary']
    elif is_compliance:
        fill_color = COLORS['accent']
    elif is_branch:
        fill_color = RGBColor(0, 180, 220)
    elif is_final:
        fill_color = RGBColor(0, 220, 150)
    else:
        fill_color = COLORS['secondary']
    
    box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        left, top, width, height
    )
    box.fill.solid()
    box.fill.fore_color.rgb = fill_color
    box.line.width = Pt(0)
    
    title_box = slide.shapes.add_textbox(
        left, top + Inches(0.15),
        width, Inches(0.4)
    )
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLORS['text_white']
    p.alignment = PP_ALIGN.CENTER
    
    if subtitle:
        sub_box = slide.shapes.add_textbox(
            left, top + Inches(0.5),
            width, Inches(0.3)
        )
        tf = sub_box.text_frame
        p = tf.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(11)
        p.font.color.rgb = COLORS['text_light']
        p.alignment = PP_ALIGN.CENTER


def add_arrow(slide, start_x, start_y, end_x, end_y):
    arrow = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW,
        start_x, start_y - Inches(0.15),
        end_x - start_x, Inches(0.3)
    )
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = COLORS['secondary']
    arrow.line.fill.background()


def add_curved_arrow(slide, start_x, start_y, end_x, end_y):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.ELBOW,
        start_x, start_y, end_x, end_y
    )
    connector.line.width = Pt(3)
    connector.line.color.rgb = COLORS['secondary']


def add_arrow_down(slide, start_x, start_y, end_x, end_y):
    arrow = slide.shapes.add_shape(
        MSO_SHAPE.DOWN_ARROW,
        start_x - Inches(0.15), start_y,
        Inches(0.3), end_y - start_y
    )
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = COLORS['secondary']
    arrow.line.fill.background()


def add_core_capability(slide):
    cap_bg = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(8.5), Inches(4.5),
        Inches(4), Inches(2.5)
    )
    cap_bg.fill.solid()
    cap_bg.fill.fore_color.rgb = COLORS['card_bg']
    cap_bg.line.color.rgb = COLORS['accent']
    cap_bg.line.width = Pt(2)
    cap_bg.fill.transparency = 0.3
    
    title_box = slide.shapes.add_textbox(
        Inches(8.7), Inches(4.7),
        Inches(3.6), Inches(0.5)
    )
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "★ 核心能力"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLORS['accent']
    
    content = [
        "• 全链路咨询服务",
        "• 一站式解决方案",
        "• 专业团队支持",
        "• 成功案例丰富"
    ]
    
    content_box = slide.shapes.add_textbox(
        Inches(8.7), Inches(5.2),
        Inches(3.6), Inches(1.6)
    )
    tf = content_box.text_frame
    tf.word_wrap = True
    
    for i, item in enumerate(content):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(13)
        p.font.color.rgb = COLORS['text_light']
        p.space_after = Pt(6)


def add_footer(slide, prs):
    page_box = slide.shapes.add_textbox(
        prs.slide_width - Inches(1.8), prs.slide_height - Inches(0.4),
        Inches(1.5), Inches(0.3)
    )
    tf = page_box.text_frame
    p = tf.paragraphs[0]
    p.text = "01"
    p.font.size = Pt(14)
    p.font.color.rgb = COLORS['text_gray']
    p.alignment = PP_ALIGN.RIGHT


if __name__ == '__main__':
    create_data_service_flow_ppt()
